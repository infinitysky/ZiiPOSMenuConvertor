"""
PE Menu Excel parser.
Reads the PE POS export format and converts to ZiiPOS-compatible data.

PE Excel columns (no header row):
  A: Product code / PE Item ID (used as ZiiPOS ItemCode)
  B: Category path (e.g. "居酒屋メニュー/Izakaya menu-揚げ物/Fried food")
  C: Item name
  D: Image (embedded, cell value is None)
  E: Price string (e.g. "￥480", "￥1,100")
  F: Tax info (e.g. "10%外税")
  G: Numeric code
  H: Status ("展示" = show / "隐藏" = hide)
  I: (empty)
"""

import os
import re
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage


def _parse_price(price_str: str) -> float:
    """Parse price string like '￥1,100' or '¥480' to float."""
    if not price_str:
        return 0.0
    cleaned = re.sub(r'[¥￥\s,，]', '', str(price_str))
    try:
        return float(cleaned)
    except (ValueError, TypeError):
        return 0.0


def _parse_tax_rate(tax_str: str) -> float:
    """Parse tax string like '10%外税' to numeric rate (10)."""
    if not tax_str:
        return 10.0
    m = re.search(r'(\d+(?:\.\d+)?)\s*%', str(tax_str))
    if m:
        return float(m.group(1))
    return 10.0


def _parse_category_path(cat_path: str) -> tuple[str, str]:
    """
    Parse PE category path into (MenuGroup, Category).

    Examples:
      "居酒屋メニュー/Izakaya menu-揚げ物/Fried food"
        → ("居酒屋メニュー", "揚げ物")
      "ちょい足し/Side dish-サイドメニュー"
        → ("ちょい足し", "サイドメニュー")
      "せきりょう"
        → ("せきりょう", "せきりょう")
    """
    if not cat_path:
        return ("Default", "Default")

    cat_path = str(cat_path).strip()

    dash_idx = cat_path.find('-')
    if dash_idx < 0:
        jp_part = cat_path.split('/')[0].strip()
        return (jp_part or "Default", jp_part or "Default")

    mg_part = cat_path[:dash_idx].strip()
    cat_part = cat_path[dash_idx + 1:].strip()

    mg_jp = mg_part.split('/')[0].strip()
    cat_jp = cat_part.split('/')[0].strip()

    return (mg_jp or "Default", cat_jp or "Default")


def _has_header_row(ws) -> bool:
    """
    Heuristic: check if the first row looks like a header.
    PE exports typically have no header; row 1 starts with an integer ID.
    """
    first_cell = ws.cell(row=1, column=1).value
    if first_cell is None:
        return False
    if isinstance(first_cell, (int, float)):
        return False
    header_keywords = {
        'id', 'code', 'itemcode', 'item', 'name', 'category',
        'price', 'tax', 'status', '商品', '名称', '分类', '价格',
    }
    first_str = str(first_cell).strip().lower()
    return first_str in header_keywords


def _sanitize_filename(name: str) -> str:
    """Make a string safe for use as a filename."""
    name = str(name).strip()
    name = re.sub(r'[\u3000]+', ' ', name)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name)
    name = name.strip('. ')
    if not name:
        name = "unnamed"
    if len(name) > 120:
        name = name[:120]
    return name


def pe_item_code(item: dict, idx: int) -> str:
    """Use PE column A (product code) when present, else sequential fallback."""
    pe_id = str(item.get("pe_id", "")).strip()
    if pe_id:
        return pe_id
    return "%04d" % (idx + 1)


def read_pe_menu(filepath: str) -> list[dict]:
    """
    Read a PE Menu Excel file and return structured menu items.

    Returns list of dicts:
      {pe_id, menu_group, category, name, price, tax_rate, status, row_idx}
      pe_id = column A product code
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    start_row = 2 if _has_header_row(ws) else 1

    items = []
    for row_idx in range(start_row, ws.max_row + 1):
        col_a = ws.cell(row=row_idx, column=1).value
        col_b = ws.cell(row=row_idx, column=2).value
        col_c = ws.cell(row=row_idx, column=3).value
        col_e = ws.cell(row=row_idx, column=5).value
        col_f = ws.cell(row=row_idx, column=6).value
        col_h = ws.cell(row=row_idx, column=8).value

        name = str(col_c).strip() if col_c else ""
        if not name:
            continue

        menu_group, category = _parse_category_path(col_b)
        price = _parse_price(col_e)
        tax_rate = _parse_tax_rate(col_f)

        status_raw = str(col_h).strip() if col_h else ""
        status = "show" if status_raw in ("展示", "表示", "show", "Show") else "hide"

        items.append({
            "pe_id": str(col_a) if col_a else "",
            "menu_group": menu_group,
            "category": category,
            "name": name,
            "price": price,
            "tax_rate": tax_rate,
            "status": status,
            "row_idx": row_idx,
        })

    wb.close()
    return items


def extract_images(filepath: str, items: list[dict], output_dir: str) -> int:
    """
    Extract embedded images from the PE Excel and save to output_dir/pics/.
    Image filenames match the item name on the same row.

    Returns the number of images saved.
    """
    pics_dir = os.path.join(output_dir, "pics")
    os.makedirs(pics_dir, exist_ok=True)

    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    row_to_name = {}
    for item in items:
        row_to_name[item["row_idx"] - 1] = item["name"]

    saved_count = 0
    name_counter = {}

    for img in ws._images:
        anchor = img.anchor
        if not hasattr(anchor, '_from'):
            continue

        img_row = anchor._from.row

        item_name = row_to_name.get(img_row, "")
        if not item_name:
            item_name = f"row_{img_row + 1}"

        safe_name = _sanitize_filename(item_name)

        if safe_name in name_counter:
            name_counter[safe_name] += 1
            safe_name = f"{safe_name}_{name_counter[safe_name]}"
        else:
            name_counter[safe_name] = 0

        try:
            from io import BytesIO
            from PIL import Image as PILImage

            img_data = img._data()
            pil_img = PILImage.open(BytesIO(img_data))

            ext = ".png"
            if pil_img.format:
                ext = f".{pil_img.format.lower()}"
                if ext == ".jpeg":
                    ext = ".jpg"

            out_path = os.path.join(pics_dir, f"{safe_name}{ext}")
            pil_img.save(out_path)
            saved_count += 1
        except Exception as e:
            print(f"[WARN] Failed to save image for row {img_row + 1} ({item_name}): {e}")

    wb.close()
    return saved_count
